# -*- mode:python; coding:utf-8 -*-

# Copyright (c) 2024 IBM Corp. All rights reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""
Testing fedramp transform command functionality.

This validate high level functionality of the transform command.
"""

import argparse
import logging
import pathlib
from typing import Any, Tuple

from docx import Document  # type: ignore
from docx.document import Document as DocxDocument  # type: ignore

from tests.test_utils import verify_responses

from trestle_fedramp.commands.transform import SSPTransformCmd

test_docx_file = 'test.docx'
example_control = 'AC-1 What is the solution and how is it implemented?'


def test_transform_ssp_level_high(
    tmp_path: pathlib.Path,
    tmp_trestle_dir_with_ssp: Tuple[pathlib.Path, str],
) -> None:
    """Test Fedramp SSP transform command with FedRAMP High Baseline template."""
    tmp_trestle_dir, ssp_name = tmp_trestle_dir_with_ssp
    tmp_file = tmp_path / test_docx_file
    args = argparse.Namespace(
        ssp_name=ssp_name,
        level='high',
        output_file=str(tmp_file),
        trestle_root=tmp_trestle_dir,
        verbose=0,
        components=''
    )
    rc = SSPTransformCmd()._run(args)
    assert rc == 0

    assert tmp_file.exists()

    # Verify example control
    expected_data = {
        'a': (
            'Part a:\nThis System: Describe how Part a is satisfied within the system.\n'
            '\n[EXAMPLE]Policies: Describe how this policy component satisfies part a.\n'
        ),
        'b': (
            'Part b:\nThis System: Describe how Part b is satisfied within the system for a component.\n'
            '\n[EXAMPLE]Procedures: Describe how Part b is satisfied within the system for another component.\n'
        ),
        'c': 'Part c:'
    }

    temp_doc_output: DocxDocument = Document(str(tmp_file))
    for table in temp_doc_output.tables:
        if example_control in table.cell(0, 0).text:
            verify_responses(table, expected_data)


def test_transform_ssp_level_moderate(
    tmp_path: pathlib.Path,
    tmp_trestle_dir_with_ssp: Tuple[pathlib.Path, str],
) -> None:
    """
    Test Fedramp SSP transform command with FedRAMP Moderate Baseline template.

    Notes: The moderate is encompassed by the high so we only need to test the details of the
    high.
    """
    tmp_trestle_dir, ssp_name = tmp_trestle_dir_with_ssp
    tmp_file = tmp_path / test_docx_file
    args = argparse.Namespace(
        ssp_name=ssp_name,
        level='moderate',
        output_file=str(tmp_file),
        trestle_root=tmp_trestle_dir,
        verbose=0,
        components=''
    )
    rc = SSPTransformCmd()._run(args)
    assert rc == 0

    assert tmp_file.exists()


def test_transform_ssp_level_low(
    tmp_path: pathlib.Path,
    tmp_trestle_dir_with_ssp: Tuple[pathlib.Path, str],
) -> None:
    """
    Test Fedramp SSP transform command with FedRAMP Low Baseline template.

    Notes: The low is encompassed by the high so we only need to test the details of the
    high.
    """
    tmp_trestle_dir, ssp_name = tmp_trestle_dir_with_ssp
    tmp_file = tmp_path / test_docx_file
    args = argparse.Namespace(
        ssp_name=ssp_name,
        level='low',
        output_file=str(tmp_file),
        trestle_root=tmp_trestle_dir,
        verbose=0,
        components=''
    )
    rc = SSPTransformCmd()._run(args)
    assert rc == 0

    assert tmp_file.exists()


def test_transform_ssp_invalid_level(tmp_path: pathlib.Path, tmp_trestle_dir: pathlib.Path, caplog: Any) -> None:
    """Test fails with an invalid level."""
    args = argparse.Namespace(
        ssp_name='test-ssp',
        level='fake',
        output_file=str(tmp_path),
        trestle_root=tmp_trestle_dir,
        verbose=0,
        components=''
    )
    rc = SSPTransformCmd()._run(args)
    assert rc == 1


def test_transform_missing_ssp(tmp_path: pathlib.Path, tmp_trestle_dir: pathlib.Path, caplog: Any) -> None:
    """Test fails with a missing ssp."""
    args = argparse.Namespace(
        ssp_name='test-ssp',
        level='high',
        output_file=str(tmp_path),
        trestle_root=tmp_trestle_dir,
        verbose=0,
        components=''
    )
    rc = SSPTransformCmd()._run(args)
    assert rc == 1

    assert any(
        record.levelno == logging.ERROR
        and 'Input ssp test-ssp does not exist in the trestle workspace.' in record.message for record in caplog.records
    )
